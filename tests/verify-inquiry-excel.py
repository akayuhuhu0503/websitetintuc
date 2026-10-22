import sys,json
from pathlib import Path
from openpyxl import load_workbook
folder=Path(sys.argv[1]);pii=load_workbook(folder/'pii-test.xlsx',read_only=True,data_only=False);minimal=load_workbook(folder/'minimal-test.xlsx',read_only=True,data_only=False)
sheet=pii.active;cells=list(sheet.iter_rows());minrows=list(minimal.active.iter_rows())
assert len(cells)>1 and len(cells[0])==11
assert len(minrows[0])==6
assert not any(c.data_type=='f' for row in cells for c in row)
assert any(str(c.value).startswith('=HYPERLINK') and c.data_type=='s' for row in cells for c in row)
assert any('tiếng Việt' in str(c.value) for row in cells for c in row)
assert any('🗺️' in str(c.value) for row in cells for c in row)
assert all(c.data_type=='s' for row in cells for c in row if c.value is not None)
assert not any('HYPERLINK' in str(c.value) or '@landportal.invalid' in str(c.value) for row in minrows for c in row)
boundary=load_workbook(folder/'formula-boundaries.xlsx',read_only=True,data_only=False)
boundarycells=[row[0] for row in boundary.active.iter_rows()]
assert [c.value for c in boundarycells]==['=SUM(1,2)','+SUM(1,2)','-1+2','@DDE','\t=HYPERLINK("x")','Tiếng Việt 🗺️']
assert all(c.data_type=='s' for c in boundarycells)
report={'library':'openpyxl','rowsRead':len(cells)-1,'fieldsPii':11,'fieldsMinimal':6,'formulas':0,'textCellsVerified':sum(c.value is not None for row in cells for c in row),'unicode':True,'unicodeSupplementary':True,'formulaBoundaryCells':len(boundarycells),'status':'PASS'}
(folder/'excel-readback.json').write_text(json.dumps(report,indent=2),encoding='utf-8');print(json.dumps(report))
